import sys

import openpyxl

sys.path.insert(0, __file__.rsplit('/', 1)[0])
from build_dashboard import MONTH_SHEETS

WORKERS_HEADER_ROW = 2
WORKERS_FIRST_ROW = 3
WORKERS_LAST_ROW = 52  # 50-worker capacity

PAYENTRIES_HEADER_ROW = 2
PAYENTRIES_FIRST_ROW = 3
PAYENTRIES_LAST_ROW = 502  # 500-entry capacity


def _add_workers_sheet(ws):
    ws.cell(row=1, column=1, value='Workers')
    headers = ['Worker ID', 'Name', 'Role', 'Rate Type', 'Rate or Agreed Amount',
               'Default Payment Mode', 'Active', 'Active Name']
    for i, header in enumerate(headers, start=1):
        ws.cell(row=WORKERS_HEADER_ROW, column=i, value=header)

    for row in range(WORKERS_FIRST_ROW, WORKERS_LAST_ROW + 1):
        ws.cell(row=row, column=8, value='=IF(G%d="Y",B%d,"")' % (row, row))


def _add_payentries_sheet(ws):
    ws.cell(row=1, column=1, value='PayEntries')
    headers = ['Worker ID', 'Worker Name', 'Month', 'Days Worked', 'Gross Pay',
               'Advance/Deduction', 'Net Pay', 'Payment Mode', 'Payments Row Ref',
               'Comments', 'Match Key']
    for i, header in enumerate(headers, start=1):
        ws.cell(row=PAYENTRIES_HEADER_ROW, column=i, value=header)

    for row in range(PAYENTRIES_FIRST_ROW, PAYENTRIES_LAST_ROW + 1):
        worker_id_cell = 'A%d' % row
        rate_type_lookup = 'INDEX(Workers!$D$%d:$D$%d,MATCH(%s,Workers!$A$%d:$A$%d,0))' % (
            WORKERS_FIRST_ROW, WORKERS_LAST_ROW, worker_id_cell, WORKERS_FIRST_ROW, WORKERS_LAST_ROW)
        rate_lookup = 'INDEX(Workers!$E$%d:$E$%d,MATCH(%s,Workers!$A$%d:$A$%d,0))' % (
            WORKERS_FIRST_ROW, WORKERS_LAST_ROW, worker_id_cell, WORKERS_FIRST_ROW, WORKERS_LAST_ROW)
        name_lookup = 'INDEX(Workers!$B$%d:$B$%d,MATCH(%s,Workers!$A$%d:$A$%d,0))' % (
            WORKERS_FIRST_ROW, WORKERS_LAST_ROW, worker_id_cell, WORKERS_FIRST_ROW, WORKERS_LAST_ROW)
        default_mode_lookup = 'INDEX(Workers!$F$%d:$F$%d,MATCH(%s,Workers!$A$%d:$A$%d,0))' % (
            WORKERS_FIRST_ROW, WORKERS_LAST_ROW, worker_id_cell, WORKERS_FIRST_ROW, WORKERS_LAST_ROW)

        ws.cell(row=row, column=2, value='=IF(%s="","",IFERROR(%s,"Invalid Worker ID"))' % (worker_id_cell, name_lookup))
        ws.cell(row=row, column=5, value=(
            '=IF(%s="","",IFERROR(IF(%s="Per Day",D%d*%s,%s),"Invalid Worker ID"))'
            % (worker_id_cell, rate_type_lookup, row, rate_lookup, rate_lookup)
        ))
        ws.cell(row=row, column=7, value='=IF(OR(%s="",E%d="Invalid Worker ID"),"",E%d-F%d)' % (worker_id_cell, row, row, row))
        ws.cell(row=row, column=8, value='=IF(%s="","",IFERROR(%s,"Invalid Worker ID"))' % (worker_id_cell, default_mode_lookup))
        ws.cell(row=row, column=11, value='=A%d&"|"&C%d' % (row, row))

    ws.column_dimensions['K'].hidden = True


def build(path):
    wb = openpyxl.load_workbook(path, keep_vba=True)

    if 'Workers' in wb.sheetnames:
        del wb['Workers']
    workers_ws = wb.create_sheet('Workers')
    _add_workers_sheet(workers_ws)

    if 'PayEntries' in wb.sheetnames:
        del wb['PayEntries']
    payentries_ws = wb.create_sheet('PayEntries')
    _add_payentries_sheet(payentries_ws)

    wb.save(path)


if __name__ == '__main__':
    build(sys.argv[1] if len(sys.argv) > 1 else 'expense_tracker_2026.xlsm')
