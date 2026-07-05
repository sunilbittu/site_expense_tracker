import shutil
import os
import openpyxl

SOURCE_WORKBOOK = os.path.join(os.path.dirname(__file__), '..', 'expense_tracker_2026.xlsx')


def sample_workbook(tmp_dir):
    """Copy the real workbook into tmp_dir and inject one sample Receipts
    row and one sample Payments row into Jan2026. Returns the new path."""
    dest = os.path.join(tmp_dir, 'sample.xlsx')
    shutil.copyfile(SOURCE_WORKBOOK, dest)

    wb = openpyxl.load_workbook(dest)
    ws = wb['Jan2026']
    # Receipts row: SNO(A) Date(B) Receipts(C) Amount(D) Category(E) Payment Mode(F)
    ws['A3'] = 1
    ws['D3'] = 1000
    ws['E3'] = 'Sales'
    ws['F3'] = 'Bank'
    # Payments row: Date(I) Payments(J) Amount(K) Main Category(L) Sub-Category(M) Payment Mode(N)
    ws['K3'] = 400
    ws['L3'] = 'Materials'
    ws['M3'] = 'Cement'
    ws['N3'] = 'Cash'
    wb.save(dest)
    return dest
