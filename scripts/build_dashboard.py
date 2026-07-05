import sys

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

MONTH_SHEETS = [
    'Jan2026', 'Feb2026', 'Mar2026', 'Apr2026', 'May2026', 'Jun2026',
    'Jul2026', 'Aug2026', 'Sep2026', 'Oct2026', 'Nov2026', 'Dec2026',
]

MONTHLY_TREND_HEADER_ROW = 44
MONTHLY_TREND_FIRST_ROW = 45
MONTHLY_TREND_LAST_ROW = MONTHLY_TREND_FIRST_ROW + len(MONTH_SHEETS) - 1  # 56

KPI_LABEL_ROW = 3
KPI_VALUE_ROW = 4
KPI_CELLS = {'balance': 'A4', 'receipts': 'C4', 'payments': 'E4', 'net': 'G4'}

MAIN_CAT_HEADER_ROW = 44
MAIN_CAT_FIRST_ROW = 45
MAIN_CAT_LAST_ROW = 54  # 10 main categories

INCOME_CAT_HEADER_ROW = 44
INCOME_CAT_FIRST_ROW = 45
INCOME_CAT_LAST_ROW = 48  # 4 income categories

BANKCASH_HEADER_ROW = 44
BANKCASH_FIRST_ROW = 45
BANKCASH_LAST_ROW = 46  # Bank, Cash

SUBCAT_HEADER_ROW = 60
SUBCAT_FIRST_ROW = 61
SUBCAT_LAST_ROW = 128  # 68 sub-categories across all main categories


def read_categories(wb):
    """Read Income Categories, Main Categories, and the Main->Sub-category
    map directly from the existing Categories sheet."""
    cat = wb['Categories']

    income_categories = []
    r = 3
    while cat.cell(row=r, column=1).value:
        income_categories.append(cat.cell(row=r, column=1).value)
        r += 1

    main_categories = []
    r = 3
    while cat.cell(row=r, column=6).value:
        main_categories.append(cat.cell(row=r, column=6).value)
        r += 1

    main_sub = {}
    col = 8
    while cat.cell(row=2, column=col).value:
        main = cat.cell(row=2, column=col).value
        subs = []
        r = 3
        while cat.cell(row=r, column=col).value:
            subs.append(cat.cell(row=r, column=col).value)
            r += 1
        main_sub[main] = subs
        col += 1

    return income_categories, main_categories, main_sub


def _add_monthly_trend_table(ws):
    ws.cell(row=MONTHLY_TREND_HEADER_ROW, column=1, value='Month')
    ws.cell(row=MONTHLY_TREND_HEADER_ROW, column=2, value='Receipts')
    ws.cell(row=MONTHLY_TREND_HEADER_ROW, column=3, value='Payments')
    ws.cell(row=MONTHLY_TREND_HEADER_ROW, column=4, value='Closing Balance')

    for i, month in enumerate(MONTH_SHEETS):
        row = MONTHLY_TREND_FIRST_ROW + i
        ws.cell(row=row, column=1, value=month)
        ws.cell(row=row, column=2, value="='%s'!B104" % month)
        ws.cell(row=row, column=3, value="='%s'!B105" % month)
        ws.cell(row=row, column=4, value="='%s'!B106" % month)


def _add_kpi_cards(ws):
    trend_receipts_range = "B%d:B%d" % (MONTHLY_TREND_FIRST_ROW, MONTHLY_TREND_LAST_ROW)
    trend_payments_range = "C%d:C%d" % (MONTHLY_TREND_FIRST_ROW, MONTHLY_TREND_LAST_ROW)

    ws['A3'] = 'Current Cash Balance'
    ws['A4'] = "='%s'!B106" % MONTH_SHEETS[-1]

    ws['C3'] = 'YTD Total Receipts'
    ws['C4'] = "=SUM(%s)" % trend_receipts_range

    ws['E3'] = 'YTD Total Payments'
    ws['E4'] = "=SUM(%s)" % trend_payments_range

    ws['G3'] = 'Net Position (YTD)'
    ws['G4'] = "=%s-%s" % (KPI_CELLS['receipts'], KPI_CELLS['payments'])


def _main_category_formula(main_cat_cell):
    terms = ["SUMIFS(%s!K3:K100,%s!L3:L100,%s)" % (m, m, main_cat_cell) for m in MONTH_SHEETS]
    return "=" + "+".join(terms)


def _income_category_formula(income_cat_cell):
    terms = ["SUMIFS(%s!D3:D100,%s!E3:E100,%s)" % (m, m, income_cat_cell) for m in MONTH_SHEETS]
    return "=" + "+".join(terms)


def _payment_mode_formula(mode_cell):
    terms = []
    for m in MONTH_SHEETS:
        terms.append("SUMIFS(%s!D3:D100,%s!F3:F100,%s)" % (m, m, mode_cell))
        terms.append("SUMIFS(%s!K3:K100,%s!N3:N100,%s)" % (m, m, mode_cell))
    return "=" + "+".join(terms)


def _add_category_tables(ws, income_categories, main_categories):
    ws.cell(row=MAIN_CAT_HEADER_ROW, column=6, value='Main Category')
    ws.cell(row=MAIN_CAT_HEADER_ROW, column=7, value='Amount')
    for i, cat in enumerate(main_categories):
        row = MAIN_CAT_FIRST_ROW + i
        ws.cell(row=row, column=6, value=cat)
        ws.cell(row=row, column=7, value=_main_category_formula("$F%d" % row))

    ws.cell(row=INCOME_CAT_HEADER_ROW, column=9, value='Income Category')
    ws.cell(row=INCOME_CAT_HEADER_ROW, column=10, value='Amount')
    for i, cat in enumerate(income_categories):
        row = INCOME_CAT_FIRST_ROW + i
        ws.cell(row=row, column=9, value=cat)
        ws.cell(row=row, column=10, value=_income_category_formula("$I%d" % row))

    ws.cell(row=BANKCASH_HEADER_ROW, column=12, value='Payment Mode')
    ws.cell(row=BANKCASH_HEADER_ROW, column=13, value='Amount')
    for i, mode in enumerate(['Bank', 'Cash']):
        row = BANKCASH_FIRST_ROW + i
        ws.cell(row=row, column=12, value=mode)
        ws.cell(row=row, column=13, value=_payment_mode_formula("$L%d" % row))


def _subcategory_formula(main_cell, sub_cell):
    terms = [
        "SUMIFS(%s!K3:K100,%s!L3:L100,%s,%s!M3:M100,%s)" % (m, m, main_cell, m, sub_cell)
        for m in MONTH_SHEETS
    ]
    return "=" + "+".join(terms)


def _add_subcategory_detail_table(ws, main_sub):
    ws.cell(row=SUBCAT_HEADER_ROW, column=1, value='Main Category')
    ws.cell(row=SUBCAT_HEADER_ROW, column=2, value='Sub-Category')
    ws.cell(row=SUBCAT_HEADER_ROW, column=3, value='Amount')

    row = SUBCAT_FIRST_ROW
    for main, subs in main_sub.items():
        for sub in subs:
            ws.cell(row=row, column=1, value=main)
            ws.cell(row=row, column=2, value=sub)
            ws.cell(row=row, column=3, value=_subcategory_formula("$A%d" % row, "$B%d" % row))
            row += 1

    table = Table(displayName='SubCategoryDetail', ref='A%d:C%d' % (SUBCAT_HEADER_ROW, SUBCAT_LAST_ROW))
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)


def build(path):
    wb = openpyxl.load_workbook(path)
    if 'Dashboard' in wb.sheetnames:
        del wb['Dashboard']
    ws = wb.create_sheet('Dashboard', 0)
    income_categories, main_categories, main_sub = read_categories(wb)

    ws.cell(row=1, column=1, value='Site Expense Dashboard')
    _add_monthly_trend_table(ws)
    _add_kpi_cards(ws)
    _add_category_tables(ws, income_categories, main_categories)
    _add_subcategory_detail_table(ws, main_sub)

    wb.save(path)


if __name__ == '__main__':
    build(sys.argv[1] if len(sys.argv) > 1 else 'expense_tracker_2026.xlsx')
