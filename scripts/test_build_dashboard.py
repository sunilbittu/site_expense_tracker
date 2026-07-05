# scripts/test_build_dashboard.py
import os
import sys
import tempfile

import formulas

sys.path.insert(0, os.path.dirname(__file__))
import build_dashboard
from test_helpers import sample_workbook

REAL_WORKBOOK = os.path.join(os.path.dirname(__file__), '..', 'expense_tracker_2026.xlsm')


def calc(path):
    xl = formulas.ExcelModel().loads(path).finish()
    return xl.calculate()


def cell_value(sol, path, sheet, cell):
    key = "'[%s]%s'!%s" % (os.path.basename(path), sheet.upper(), cell)
    v = sol[key].value
    try:
        return v[0][0]
    except (TypeError, IndexError):
        return v


def test_read_categories():
    import openpyxl
    wb = openpyxl.load_workbook(REAL_WORKBOOK)
    income, main, main_sub = build_dashboard.read_categories(wb)
    assert income == ['Investor Capital', 'Partner Contribution', 'Loan Received', 'Sales']
    assert len(main) == 12
    assert sum(len(v) for v in main_sub.values()) == 180
    print("test_read_categories PASSED")


def test_monthly_trend_table():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_dashboard.build(path)
        sol = calc(path)
        # Jan2026 row is MONTHLY_TREND_FIRST_ROW (45): Receipts=1000, Payments=400, Closing=600
        row = build_dashboard.MONTHLY_TREND_FIRST_ROW
        assert cell_value(sol, path, 'Dashboard', 'B%d' % row) == 1000
        assert cell_value(sol, path, 'Dashboard', 'C%d' % row) == 400
        assert cell_value(sol, path, 'Dashboard', 'D%d' % row) == 600
        print("test_monthly_trend_table PASSED")


def test_kpi_cards():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_dashboard.build(path)
        sol = calc(path)
        assert cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['receipts']) == 1000
        assert cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['payments']) == 400
        assert cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['net']) == 600
        # Dec2026 closing balance carries Jan's 600 forward through empty months
        assert cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['balance']) == 600
        print("test_kpi_cards PASSED")


def test_cash_hand_and_bank_kpis():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_dashboard.build(path)
        sol = calc(path)
        # Fixture: Jan2026 has a 1000 Bank receipt and a 400 Cash payment.
        # Cash at Hand = Cash receipts (0) - Cash payments (400) = -400
        # Cash at Bank = Bank receipts (1000) - Bank payments (0) = 1000
        assert cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['cash_hand']) == -400
        assert cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['cash_bank']) == 1000
        # Cash at Hand + Cash at Bank must always equal Current Cash Balance
        balance = cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['balance'])
        cash_hand = cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['cash_hand'])
        cash_bank = cell_value(sol, path, 'Dashboard', build_dashboard.KPI_CELLS['cash_bank'])
        assert cash_hand + cash_bank == balance
        print("test_cash_hand_and_bank_kpis PASSED")


def test_category_tables():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_dashboard.build(path)
        sol = calc(path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        dash = wb['Dashboard']

        # Main Category spend: find the "Materials" row, assert amount == 400
        materials_row = None
        for r in range(build_dashboard.MAIN_CAT_FIRST_ROW, build_dashboard.MAIN_CAT_LAST_ROW + 1):
            if dash.cell(row=r, column=6).value == 'Materials':
                materials_row = r
        assert materials_row is not None
        assert cell_value(sol, path, 'Dashboard', 'G%d' % materials_row) == 400

        # Income by Category: find the "Sales" row, assert amount == 1000
        sales_row = None
        for r in range(build_dashboard.INCOME_CAT_FIRST_ROW, build_dashboard.INCOME_CAT_LAST_ROW + 1):
            if dash.cell(row=r, column=9).value == 'Sales':
                sales_row = r
        assert sales_row is not None
        assert cell_value(sol, path, 'Dashboard', 'J%d' % sales_row) == 1000

        # Bank vs Cash: Bank total == 1000 (the receipt), Cash total == 400 (the payment)
        bank_row, cash_row = build_dashboard.BANKCASH_FIRST_ROW, build_dashboard.BANKCASH_FIRST_ROW + 1
        assert dash.cell(row=bank_row, column=12).value == 'Bank'
        assert cell_value(sol, path, 'Dashboard', 'M%d' % bank_row) == 1000
        assert dash.cell(row=cash_row, column=12).value == 'Cash'
        assert cell_value(sol, path, 'Dashboard', 'M%d' % cash_row) == 400

        print("test_category_tables PASSED")


def test_subcategory_detail_table():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_dashboard.build(path)
        sol = calc(path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        dash = wb['Dashboard']

        target_row = None
        for r in range(build_dashboard.SUBCAT_FIRST_ROW, build_dashboard.SUBCAT_LAST_ROW + 1):
            if dash.cell(row=r, column=1).value == 'Materials' and dash.cell(row=r, column=2).value == 'Cement':
                target_row = r
        assert target_row is not None
        assert cell_value(sol, path, 'Dashboard', 'C%d' % target_row) == 400

        assert 'SubCategoryDetail' in dash.tables
        print("test_subcategory_detail_table PASSED")


def test_charts():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_dashboard.build(path)

        import openpyxl
        from openpyxl.chart import LineChart, BarChart, PieChart
        wb = openpyxl.load_workbook(path)
        dash = wb['Dashboard']

        assert len(dash._charts) == 4
        chart_types = [type(c) for c in dash._charts]
        assert chart_types.count(LineChart) == 1
        assert chart_types.count(BarChart) == 2
        assert chart_types.count(PieChart) == 1
        print("test_charts PASSED")


def test_builds_cleanly_against_real_workbook():
    with tempfile.TemporaryDirectory() as tmp:
        import shutil
        path = os.path.join(tmp, 'real_copy.xlsm')
        shutil.copyfile(REAL_WORKBOOK, path)
        build_dashboard.build(path)
        sol = calc(path)

        # Whatever data currently exists in the real workbook, every KPI
        # must resolve to a plain number, never an error (#REF!, #N/A, etc).
        for cell in build_dashboard.KPI_CELLS.values():
            value = cell_value(sol, path, 'Dashboard', cell)
            assert isinstance(value, (int, float)), "%s resolved to %r, expected a number" % (cell, value)

        print("test_builds_cleanly_against_real_workbook PASSED")


if __name__ == '__main__':
    test_read_categories()
    test_monthly_trend_table()
    test_kpi_cards()
    test_cash_hand_and_bank_kpis()
    test_category_tables()
    test_subcategory_detail_table()
    test_charts()
    test_builds_cleanly_against_real_workbook()
