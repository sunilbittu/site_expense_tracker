import os
import sys
import tempfile

import formulas

sys.path.insert(0, os.path.dirname(__file__))
import build_payslip
from test_helpers import sample_workbook

REAL_WORKBOOK = os.path.join(os.path.dirname(__file__), '..', 'expense_tracker_2026.xlsm')


def calc(path):
    xl = formulas.ExcelModel().loads(path).finish()
    return xl.calculate()


def cell_value(sol, path, sheet, cell):
    key = "'[%s]%s'!%s" % (os.path.basename(path), sheet.upper(), cell)
    v = sol[key].value
    try:
        return v[0][0]
    except (TypeError, IndexError):
        return v


def test_workers_sheet_structure_and_active_name_formula():
    with tempfile.TemporaryDirectory() as tmp:
        path = sample_workbook(tmp)
        build_payslip.build(path)

        import openpyxl
        wb = openpyxl.load_workbook(path)
        workers = wb['Workers']
        assert workers.cell(row=build_payslip.WORKERS_HEADER_ROW, column=1).value == 'Worker ID'
        assert workers.cell(row=build_payslip.WORKERS_HEADER_ROW, column=8).value == 'Active Name'

        # Row 3: an Active labourer -> Active Name should mirror Name
        workers.cell(row=3, column=1, value='W001')
        workers.cell(row=3, column=2, value='Ramesh')
        workers.cell(row=3, column=3, value='Labourer')
        workers.cell(row=3, column=4, value='Per Day')
        workers.cell(row=3, column=5, value=500)
        workers.cell(row=3, column=6, value='Cash')
        workers.cell(row=3, column=7, value='Y')
        # Row 4: an Inactive worker -> Active Name should be blank
        workers.cell(row=4, column=1, value='W002')
        workers.cell(row=4, column=2, value='Suresh')
        workers.cell(row=4, column=3, value='Labourer')
        workers.cell(row=4, column=4, value='Per Day')
        workers.cell(row=4, column=5, value=500)
        workers.cell(row=4, column=6, value='Cash')
        workers.cell(row=4, column=7, value='N')
        wb.save(path)

        sol = calc(path)
        assert cell_value(sol, path, 'Workers', 'H3') == 'Ramesh'
        assert cell_value(sol, path, 'Workers', 'H4') == ''
        print("test_workers_sheet_structure_and_active_name_formula PASSED")


if __name__ == '__main__':
    test_workers_sheet_structure_and_active_name_formula()
